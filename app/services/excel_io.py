"""Excel export and import. Import reads sheets by header name so the original Excel_Native_ATS.xlsx loads directly."""
import io
from datetime import date, datetime, time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..fields import STATUSES
from . import audit, candidates as cand_svc, interviews as int_svc
from .common import Ctx, Forbidden, ValidationError
from .dashboard import recruiter_report
from .validation import clean_candidate, parse_ref

NAVY, BLUE = "1F3864", "2F6DB5"

MASTER_COLUMNS = [
    ("Candidate ID", "candidate_id"), ("Date Added", "date_added"), ("Candidate Name", "full_name"),
    ("Current Company", "current_company"), ("Location", "current_location"),
    ("Experience (Yrs)", "total_experience_years"), ("Current CTC (LPA)", "current_ctc_lakhs"),
    ("Expected CTC (LPA)", "expected_ctc_lakhs"), ("Notice Period (Days)", "official_notice_period"),
    ("Earliest Joining Date", "earliest_joining_date"), ("How Soon Can Join", "how_soon_can_join"),
    ("Existing Offer?", "existing_offer"), ("Offer Details", "offer_details"),
    ("Date of Birth (Optional)", "date_of_birth"), ("Email", "email"), ("Phone", "phone"),
    ("Recruiter", "recruiter"), ("Client", "client"), ("Role", "primary_skill_role"),
    ("Submitted to Client?", "profile_submitted_to_client"), ("Client Submission Date", "client_submission_date"),
    ("Current Status", "current_status"), ("Duplicate Check", "duplicate_check"),
    ("Job Req ID", "job_req_id"), ("Recruiter Notes", "recruiter_notes"),
]
TRACK_COLUMNS = [
    ("Sr No", None), ("Candidate ID", "candidate_id"), ("Candidate Name", "full_name"), ("Client", "client"),
    ("Interview Stage", "interview_stage"), ("Interview Date", "interview_date"), ("Time", "interview_time"),
    ("Panel", "interviewer_panel"), ("Feedback", "feedback"), ("Status", "status"),
    ("Next Action", "next_action"), ("Follow-up Date", "next_followup_date"),
]
REPORT_COLUMNS = [("Recruiter", "recruiter"), ("Candidates Added", "candidates_added"),
                  ("Client Submissions", "client_submissions"), ("Offers", "offers"),
                  ("Drops", "drops"), ("Yet to Join", "yet_to_join")]
TRACK_KEYS = {"Interview Stage": "interview_stage", "Interview Date": "interview_date", "Time": "interview_time",
              "Panel": "interviewer_panel", "Feedback": "feedback", "Status": "status",
              "Next Action": "next_action", "Follow-up Date": "next_followup_date"}
MASTER_KEYS = {h: k for h, k in MASTER_COLUMNS if k not in ("candidate_id", "duplicate_check")}


def _sheet(wb, title, heading, columns, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws["A1"] = heading
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    for c in range(1, len(columns) + 1):
        ws.cell(1, c).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 26
    for i, (h, _) in enumerate(columns, 1):
        c = ws.cell(4, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(32, len(h) + 4))
    for r, row in enumerate(rows, 5):
        for i, v in enumerate(row, 1):
            cell = ws.cell(r, i, v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"  # user text must never be stored as a formula
    ws.freeze_panes = "A5"
    if rows:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{4 + len(rows)}"


def export_workbook(ctx: Ctx) -> bytes:
    cands = cand_svc.list_candidates(ctx, limit=None)["items"]
    interviews = int_svc.list_interviews(ctx)
    wb = Workbook()
    master = list(MASTER_COLUMNS) + ([("Vendor", "vendor")] if not ctx.vendor_id and any(c.get("vendor") for c in cands) else [])
    _sheet(wb, "Candidate Master", "Candidate Master | Central Database", master,
           [[c.get(k) for _, k in master] for c in cands], first=True)
    _sheet(wb, "Interview Tracker", "Interview Tracker | Pipeline", TRACK_COLUMNS,
           [[n] + [i.get(k) for _, k in TRACK_COLUMNS[1:]] for n, i in enumerate(interviews, 1)])
    _sheet(wb, "Recruiter Reports", "Recruiter Reports | Recruiter-wise Summary", REPORT_COLUMNS,
           [[r[k] for _, k in REPORT_COLUMNS] for r in recruiter_report(ctx)])
    buf = io.BytesIO()
    wb.save(buf)
    audit.record(ctx, "export", "candidate", None, {"rows": len(cands)})
    ctx.db.commit()
    return buf.getvalue()


def _text(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    return "" if s.startswith("=") else s


def _find_header(ws, must_have):
    for r in range(1, 11):
        vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)]
        if must_have in vals:
            return r, {v: i + 1 for i, v in enumerate(vals) if v}
    return None, {}


def import_workbook(ctx: Ctx, blob: bytes, row_cap: int = 5000) -> dict:
    if not ctx.can_write:
        raise Forbidden("your role is read-only")
    try:
        wb = load_workbook(io.BytesIO(blob), data_only=False)
    except Exception:
        raise ValueError("This file is not a readable .xlsx workbook.")
    if "Candidate Master" not in wb.sheetnames:
        raise ValueError("No 'Candidate Master' sheet found in this workbook.")
    ws = wb["Candidate Master"]
    hr, cols = _find_header(ws, "Candidate Name")
    if not hr:
        raise ValueError("Could not find the 'Candidate Name' header on Candidate Master.")
    if ws.max_row - hr > row_cap:
        raise ValueError(f"Too many rows: the limit is {row_cap} per import.")
    
    summary = {"added": 0, "skipped_blank": 0, "skipped_duplicate": 0, "added_with_name_flag": 0,
               "invalid": [], "interviews_added": 0}
    id_map = {}  # workbook candidate number -> newly added candidate ref
    today = ctx.today().isoformat()
    try:
        for r in range(hr + 1, ws.max_row + 1):
            raw = {h: _text(ws.cell(r, c).value) for h, c in cols.items()}
            if not raw.get("Candidate Name"):
                summary["skipped_blank"] += 1
                continue
            data = {}
            for h, key in MASTER_KEYS.items():
                v = raw.get(h, "")
                if h == "Notice Period (Days)" and v.isdigit():
                    v = "Immediate" if v == "0" else f"{v} Days"
                if v:
                    data[key] = v
            if data.get("current_status") not in STATUSES:
                data.pop("current_status", None)
            old = parse_ref(raw.get("Candidate ID")) or (r - hr)
            try:
                vals = clean_candidate(data, partial=True)
            except ValidationError as e:
                summary["invalid"].append({"row": r, "errors": e.errors})
                continue
            vals["date_added"] = vals.get("date_added") or today
            matches = cand_svc.find_matches(ctx, vals["full_name"], vals.get("email") or "", vals.get("phone") or "")
            if any({"email", "phone"} & set(m["match"]) and "name" in m["match"] for m in matches):
                summary["skipped_duplicate"] += 1
                continue
            c = cand_svc.insert_cleaned(ctx, vals)
            summary["added"] += 1
            if matches:
                summary["added_with_name_flag"] += 1
            id_map[old] = cand_svc.public_id(c.seq)

        if "Interview Tracker" in wb.sheetnames and id_map:
            ws = wb["Interview Tracker"]
            hr, cols = _find_header(ws, "Candidate ID")
            for r in range(hr + 1, ws.max_row + 1) if hr else []:
                raw = {h: _text(ws.cell(r, c).value) for h, c in cols.items()}
                ref = id_map.get(parse_ref(raw.get("Candidate ID")))
                if not ref or not raw.get("Status"):
                    continue
                data = {k: raw.get(h, "") for h, k in TRACK_KEYS.items()}
                data["candidate_id"] = ref
                try:
                    int_svc.create_interview(ctx, data, commit=False)
                    summary["interviews_added"] += 1
                except ValidationError:
                    continue
        audit.record(ctx, "import", "candidate", None, {k: v for k, v in summary.items() if k != "invalid"})
        ctx.db.commit()
    except Exception:
        ctx.db.rollback()
        raise
    return summary