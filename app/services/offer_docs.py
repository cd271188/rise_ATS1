"""Offer letter documents: find and fill {{placeholders}} in a company's Word or PDF template and produce the PDF.
No database access here; bytes in, bytes out."""
import io
import re
import shutil
import subprocess
import tempfile
import threading
import zipfile
from functools import lru_cache
from pathlib import Path
from xml.sax.saxutils import escape

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")
MAX_TEMPLATE_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED_BYTES = 30 * 1024 * 1024
MAX_PDF_PAGES = 10
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_SOFFICE_LOCK = threading.Lock()


# ---------------------------------------------------------------- detection and validation
def detect_kind(blob: bytes) -> str:
    if blob[:5] == b"%PDF-":
        return "pdf"
    if blob[:2] == b"PK":
        return "docx"
    raise ValueError("upload a Word (.docx) or PDF file")


def validate_template(blob: bytes) -> str:
    """Return 'docx' or 'pdf', or raise ValueError saying what is wrong with the file."""
    if not blob:
        raise ValueError("empty upload")
    if len(blob) > MAX_TEMPLATE_BYTES:
        raise ValueError("the template is larger than 2 MB")
    kind = detect_kind(blob)
    if kind == "docx":
        try:
            with zipfile.ZipFile(io.BytesIO(blob)) as z:
                if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                    raise ValueError("the Word file is too large when unpacked")
                if "word/document.xml" not in z.namelist():
                    raise ValueError("this is not a Word (.docx) document")
                _open_docx(blob)
        except zipfile.BadZipFile:
            raise ValueError("this is not a valid Word (.docx) document")
    else:
        _pdf_boxes(blob)
    return kind


def _open_docx(blob: bytes):
    import docx
    try:
        return docx.Document(io.BytesIO(blob))
    except Exception:
        raise ValueError("this Word document could not be read")


def _clean(value) -> str:
    return _CONTROL.sub("", "" if value is None else str(value))


def _scan(texts) -> dict:
    names, bad = set(), []
    for t in texts:
        names.update(m.group(1) for m in PLACEHOLDER.finditer(t))
        rest = PLACEHOLDER.sub("", t)
        if "{{" in rest or "}}" in rest:
            bad.append(t.strip()[:80])
    return {"names": sorted(names), "malformed": bad}


def placeholders(kind: str, blob: bytes) -> dict:
    """{'names': [...], 'malformed': [snippets of text with a broken {{ }}]}"""
    if kind == "docx":
        return _scan(p.text for p in _docx_paragraphs(_open_docx(blob)))
    return _scan(_pdf_boxes(blob)[1])


# ---------------------------------------------------------------- Word
def _in_container(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _in_container(cell)


def _docx_paragraphs(doc):
    yield from _in_container(doc)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if not part.is_linked_to_previous:
                yield from _in_container(part)


def _fill_paragraph(p, values: dict) -> None:
    runs = p.runs
    texts = [r.text for r in runs]
    matches = list(PLACEHOLDER.finditer("".join(texts)))
    if not matches:
        return
    bounds, pos = [], 0
    for t in texts:
        bounds.append((pos, pos + len(t)))
        pos += len(t)
    new = texts[:]
    for m in reversed(matches):
        if m.group(1) not in values:
            continue
        value, first = _clean(values[m.group(1)]), True
        for i, (a, b) in enumerate(bounds):
            if b <= m.start() or a >= m.end():
                continue
            lo, hi = max(m.start(), a) - a, min(m.end(), b) - a
            new[i] = new[i][:lo] + (value if first else "") + new[i][hi:]
            first = False
    for run, text in zip(runs, new):
        if text != run.text:
            run.text = text


def fill_docx(blob: bytes, values: dict) -> bytes:
    doc = _open_docx(blob)
    for p in _docx_paragraphs(doc):
        _fill_paragraph(p, values)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- fonts
_FONT_DIRS = [Path(__file__).resolve().parent.parent / "assets" / "fonts", Path("C:/Windows/Fonts"),
              Path("/usr/share/fonts/truetype/dejavu"), Path("/usr/share/fonts/dejavu"), Path("/usr/share/fonts/TTF"),
              Path("/System/Library/Fonts/Supplemental"), Path("/Library/Fonts"), Path.home() / "Library" / "Fonts"]
_FONT_SETS = [
    ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf"),
    ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"),
    ("Arial.ttf", "Arial Bold.ttf", "Arial Italic.ttf", "Arial Bold Italic.ttf"),
]


@lru_cache(maxsize=1)
def _fonts() -> tuple[str, str, str, str, bool]:
    """(regular, bold, italic, bold-italic, unicode) font names registered with reportlab."""
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for files in _FONT_SETS:
        for folder in _FONT_DIRS:
            paths = [folder / f for f in files]
            if all(p.exists() for p in paths):
                try:
                    names = ("Atssans", "Atssans-Bold", "Atssans-Italic", "Atssans-BoldItalic")
                    for name, path in zip(names, paths):
                        pdfmetrics.registerFont(TTFont(name, str(path)))
                    for bold, italic, name in ((0, 0, names[0]), (1, 0, names[1]), (0, 1, names[2]), (1, 1, names[3])):
                        addMapping("Atssans", bold, italic, name)
                    return (*names, True)
                except Exception:
                    continue
    return ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique", False)


def _text(s: str) -> str:
    """Built-in PDF fonts have no rupee sign; spell it out when no Unicode font is installed on this machine."""
    return s if _fonts()[4] else s.replace("₹", "Rs.")


# ---------------------------------------------------------------- PDF templates (placeholders overwritten in place)
def _pdf_boxes(blob: bytes) -> tuple[list[dict], list[str]]:
    """Every {{placeholder}} in a PDF with its page and position, plus the text of each line (to spot broken ones)."""
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams, LTChar, LTTextContainer, LTTextLine
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(blob))
        if reader.is_encrypted:
            raise ValueError("this PDF is password protected")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f"the PDF has more than {MAX_PDF_PAGES} pages")
        if any(int(p.get("/Rotate", 0) or 0) % 360 for p in reader.pages):
            raise ValueError("rotated PDF pages are not supported")
        pages = list(extract_pages(io.BytesIO(blob), laparams=LAParams()))
    except ValueError:
        raise
    except Exception:
        raise ValueError("this PDF could not be read")
    boxes, lines, chars_seen = [], [], 0
    for index, page in enumerate(pages):
        for block in page:
            if not isinstance(block, LTTextContainer):
                continue
            for line in block:
                if not isinstance(line, LTTextLine):
                    continue
                text, owners = "", []
                for ch in line:
                    t = ch.get_text()
                    text += t
                    owners.extend([ch if isinstance(ch, LTChar) else None] * len(t))
                chars_seen += sum(1 for o in owners if o)
                lines.append(text)
                for m in PLACEHOLDER.finditer(text):
                    own = [o for o in owners[m.start():m.end()] if o]
                    if not own:
                        continue
                    first = own[0]
                    boxes.append({"page": index, "name": m.group(1), "x0": min(o.x0 for o in own), "x1": max(o.x1 for o in own),
                                  "baseline": first.matrix[5], "size": float(first.size), "bold": "bold" in first.fontname.lower()})
    if not chars_seen:
        raise ValueError("this PDF has no selectable text (is it a scan?); upload a Word or text-based PDF")
    return boxes, lines


def fill_pdf(blob: bytes, values: dict) -> bytes:
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas
    boxes, _ = _pdf_boxes(blob)
    regular, bold, _, _, _ = _fonts()
    reader, writer = PdfReader(io.BytesIO(blob)), PdfWriter()
    for index, source in enumerate(reader.pages):
        page = writer.add_page(source)
        mine = [b for b in boxes if b["page"] == index and b["name"] in values]
        if mine:
            left, bottom = float(page.mediabox.left), float(page.mediabox.bottom)
            buf = io.BytesIO()
            c = canvas.Canvas(buf, pagesize=(float(page.mediabox.width) + left, float(page.mediabox.height) + bottom))
            for b in mine:
                size = max(b["size"], 4.0)
                x0, base = b["x0"] + left, b["baseline"] + bottom
                c.setFillColorRGB(1, 1, 1)
                c.rect(x0 - 1, base - 0.25 * size, b["x1"] - b["x0"] + 2, 1.15 * size, fill=1, stroke=0)
                c.setFillColorRGB(0, 0, 0)
                c.setFont(bold if b["bold"] else regular, size)
                c.drawString(x0, base, _text(_clean(values[b["name"]])))
            c.save()
            page.merge_page(PdfReader(io.BytesIO(buf.getvalue())).pages[0])
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# ---------------------------------------------------------------- Word to PDF
def soffice_path() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(p).exists():
            return p
    return None


def _libreoffice_pdf(exe: str, docx: bytes) -> bytes:
    with _SOFFICE_LOCK, tempfile.TemporaryDirectory() as d:
        src = Path(d) / "offer.docx"
        src.write_bytes(docx)
        subprocess.run([exe, "--headless", "--norestore", "--env:UserInstallation=" + (Path(d) / "profile").as_uri(),
                        "--convert-to", "pdf", "--outdir", d, str(src)], capture_output=True, timeout=60, check=True)
        return (Path(d) / "offer.pdf").read_bytes()


def _markup(p, base: float) -> str:
    out = []
    for run in p.runs:
        t = run.text
        if not t:
            continue
        s = escape(_text(t)).replace("\n", "<br/>").replace("\t", "&nbsp;" * 4)
        if run.font.size and 6 <= run.font.size.pt <= 36 and abs(run.font.size.pt - base) > 0.5:
            s = f'<font size="{run.font.size.pt:g}">{s}</font>'
        if run.bold:
            s = f'<b>{s}</b>'
        if run.italic:
            s = f'<i>{s}</i>'
        if run.underline:
            s = f'<u>{s}</u>'
        out.append(s)
    return "".join(out)


def _builtin_pdf(docx: bytes) -> bytes:
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph as DocxParagraph
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    from reportlab.platypus import Table as PdfTable, TableStyle

    doc = _open_docx(docx)
    regular, bold, _, _, _ = _fonts()
    normal = doc.styles["Normal"].font.size
    base = min(max(normal.pt if normal else 10.5, 8.0), 14.0)
    section = doc.sections[0]
    page_w = section.page_width.pt if section.page_width else 595.0
    page_h = section.page_height.pt if section.page_height else 842.0
    margins = [min(max((m.pt if m else 56.0), 36.0), 90.0) for m in
               (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin)]
    avail = page_w - margins[0] - margins[1]
    aligns = {0: 0, 1: 1, 2: 2, 3: 4}
    heading_sizes = {"Title": 20, "Heading 1": 16, "Heading 2": 14, "Heading 3": 12}

    def style_for(p) -> ParagraphStyle:
        name = p.style.name if p.style is not None else ""
        size = heading_sizes.get(name, base)
        is_heading = name in heading_sizes
        return ParagraphStyle(p, fontName=bold if is_heading else regular, fontSize=size, leading=size * 1.3,
                              alignment=aligns.get(p.alignment if p.alignment is not None else 0, 0), spaceAfter=4,
                              leftIndent=14 if name.startswith("List") else 0)

    def para(p):
        html = _markup(p, base)
        if (p.style.name if p.style is not None else "").startswith("List"):
            html = "&nbsp;&nbsp;" + html
        return Paragraph(html or "&nbsp;", style_for(p))

    cell_style = ParagraphStyle("cell", fontName=regular, fontSize=base - 1, leading=(base - 1) * 1.25)
    flow = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = DocxParagraph(child, doc)
            if child.xpath('.//w:br[@w:type="page"]'):
                flow.append(PageBreak())
            if p.text.strip() or p.runs:
                flow.append(para(p) if p.text.strip() else Spacer(1, 6))
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            rows = []
            for row in table.rows:
                cells, last = [], None
                for cell in row.cells:
                    if cell._tc is last:
                        cells.append("")
                        continue
                    last = cell._tc
                    html = "<br/>".join(_markup(cp, base) for cp in cell.paragraphs if cp.text.strip())
                    cells.append(Paragraph(html or "&nbsp;", cell_style))
                rows.append(cells)
            if rows and rows[0]:
                t = PdfTable(rows, colWidths=[avail / len(rows[0])] * len(rows[0]))
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
                flow += [t, Spacer(1, 8)]

    head = foot = ""
    if not section.header.is_linked_to_previous:
        head = "".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
    if not section.footer.is_linked_to_previous:
        foot = "".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())

    def decorate(canvas, _doc):
        canvas.saveState()
        canvas.setFont(regular, 8)
        if head:
            canvas.drawCentredString(page_w / 2, page_h - margins[2] / 2, _text(head)[:150])
        if foot:
            canvas.drawCentredString(page_w / 2, margins[3] / 2, _text(foot)[:150])
        canvas.restoreState()

    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=(page_w, page_h), leftMargin=margins[0], rightMargin=margins[1],
                      topMargin=margins[2], bottomMargin=margins[3], title="Offer letter").build(flow, onFirstPage=decorate, onLaterPages=decorate)
    return out.getvalue()


def render(kind: str, blob: bytes, values: dict) -> tuple[bytes, str]:
    """Fill the template and return (pdf bytes, engine used: libreoffice | builtin | overlay)."""
    if kind == "pdf":
        return fill_pdf(blob, values), "overlay"
    filled = fill_docx(blob, values)
    exe = soffice_path()
    if exe:
        try:
            return _libreoffice_pdf(exe, filled), "libreoffice"
        except Exception:
            pass
    return _builtin_pdf(filled), "builtin"